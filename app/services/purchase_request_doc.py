from datetime import date
from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.platypus import (
    SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer,
)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

_NAVY = "1F3A5F"
_ORANGE = "F97316"
_WHITE = "FFFFFF"
_GREY = "F5F5F5"


def _has_prices(pr: dict) -> bool:
    return pr.get("total_amount") is not None


# ---------- Excel ----------
def build_bc_workbook(pr: dict) -> BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = "Bon de commande"
    with_price = _has_prices(pr)

    headers = ["Référence", "Désignation", "Quantité"]
    if with_price:
        headers += ["Prix unitaire", "Montant"]
    ncol = len(headers)
    last = get_column_letter(ncol)

    ws.merge_cells(f"A1:{last}1")
    ws["A1"] = "GOCAB 225 — Bon de commande"
    ws["A1"].font = Font(size=15, bold=True, color=_NAVY)

    ws["A2"] = f"N° {pr['bc_number']}"
    ws["A2"].font = Font(bold=True, size=11)
    ws["A3"] = f"Fournisseur : {pr['supplier']['name']}"
    ws["A4"] = f"Date : {_fmt(pr['request_date'])}"
    if pr.get("expected_date"):
        ws["A5"] = f"Livraison souhaitée : {_fmt(pr['expected_date'])}"

    header_row = 7
    thin = Side(style="thin", color="DDDDDD")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    fill = PatternFill("solid", fgColor=_NAVY)
    for i, label in enumerate(headers, start=1):
        c = ws.cell(row=header_row, column=i, value=label)
        c.font = Font(bold=True, color=_WHITE)
        c.fill = fill
        c.alignment = Alignment(horizontal="center")
        c.border = border

    r = header_row + 1
    stripe = PatternFill("solid", fgColor=_GREY)
    total = Decimal("0")
    for idx, it in enumerate(pr["items"]):
        row_vals = [it["reference"], it["designation"], it["quantity"]]
        if with_price:
            up = it["unit_price"]
            lt = it["line_total"]
            row_vals += [
                float(up) if up is not None else "—",
                float(lt) if lt is not None else "—",
            ]
            if lt is not None:
                total += Decimal(str(lt))
        for col, val in enumerate(row_vals, start=1):
            c = ws.cell(row=r, column=col, value=val)
            c.border = border
            if idx % 2 == 1:
                c.fill = stripe
            if col == 3:
                c.number_format = "#,##0"
                c.alignment = Alignment(horizontal="right")
            if with_price and col in (4, 5):
                c.number_format = "#,##0"
                c.alignment = Alignment(horizontal="right")
        r += 1

    if with_price:
        lab = ws.cell(row=r, column=ncol - 1, value="TOTAL")
        lab.font = Font(bold=True)
        lab.alignment = Alignment(horizontal="right")
        tc = ws.cell(row=r, column=ncol, value=float(total))
        tc.font = Font(bold=True)
        tc.number_format = "#,##0"
        tc.fill = PatternFill("solid", fgColor="FFEAD5")

    if pr.get("notes"):
        ws.cell(row=r + 2, column=1, value=f"Notes : {pr['notes']}")

    widths = [16, 34, 12] + ([16, 16] if with_price else [])
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.sheet_view.showGridLines = False

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


# ---------- PDF ----------
def build_bc_pdf(pr: dict) -> BytesIO:
    buf = BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=A4,
        topMargin=20 * mm, bottomMargin=20 * mm,
        leftMargin=18 * mm, rightMargin=18 * mm,
    )
    styles = getSampleStyleSheet()
    navy = colors.HexColor(f"#{_NAVY}")
    orange = colors.HexColor(f"#{_ORANGE}")

    title_style = ParagraphStyle("t", parent=styles["Title"], textColor=navy, fontSize=18, spaceAfter=2)
    sub = ParagraphStyle("s", parent=styles["Normal"], fontSize=9, textColor=colors.grey)
    label = ParagraphStyle("l", parent=styles["Normal"], fontSize=10)

    with_price = _has_prices(pr)
    flow = []

    flow.append(Paragraph("GOCAB 225", title_style))
    flow.append(Paragraph("Bon de commande", ParagraphStyle("bc", parent=styles["Heading2"], textColor=orange)))
    flow.append(Spacer(1, 6 * mm))

    info = f"""
    <b>N° :</b> {pr['bc_number']}<br/>
    <b>Fournisseur :</b> {pr['supplier']['name']}<br/>
    <b>Date :</b> {_fmt(pr['request_date'])}<br/>
    """
    if pr.get("expected_date"):
        info += f"<b>Livraison souhaitée :</b> {_fmt(pr['expected_date'])}<br/>"
    flow.append(Paragraph(info, label))
    flow.append(Spacer(1, 6 * mm))

    if with_price:
        head = ["Référence", "Désignation", "Qté", "Prix unit.", "Montant"]
    else:
        head = ["Référence", "Désignation", "Qté"]
    data = [head]
    total = Decimal("0")
    for it in pr["items"]:
        row = [it["reference"], it["designation"], str(it["quantity"])]
        if with_price:
            up = it["unit_price"]
            lt = it["line_total"]
            row += [
                _money(up) if up is not None else "—",
                _money(lt) if lt is not None else "—",
            ]
            if lt is not None:
                total += Decimal(str(lt))
        data.append(row)
    if with_price:
        data.append(["", "", "", "TOTAL", _money(total)])

    col_widths = ([32 * mm, 70 * mm, 16 * mm, 28 * mm, 28 * mm] if with_price
                  else [40 * mm, 100 * mm, 20 * mm])
    table = Table(data, colWidths=col_widths, repeatRows=1)
    style = [
        ("BACKGROUND", (0, 0), (-1, 0), navy),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#DDDDDD")),
        ("ALIGN", (2, 0), (-1, -1), "RIGHT"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -2 if with_price else -1),
         [colors.white, colors.HexColor("#F7F7F7")]),
    ]
    if with_price:
        style.append(("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#FFEAD5")))
        style.append(("FONTNAME", (-2, -1), (-1, -1), "Helvetica-Bold"))
    table.setStyle(TableStyle(style))
    flow.append(table)

    if pr.get("notes"):
        flow.append(Spacer(1, 6 * mm))
        flow.append(Paragraph(f"<b>Notes :</b> {pr['notes']}", label))

    flow.append(Spacer(1, 14 * mm))
    flow.append(Paragraph("Signature / Cachet :", sub))

    doc.build(flow)
    buf.seek(0)
    return buf


def _money(v) -> str:
    return f"{Decimal(str(v)):,.0f}".replace(",", " ")


def _fmt(value) -> str:
    if isinstance(value, date):
        return value.strftime("%d/%m/%Y")
    return str(value)