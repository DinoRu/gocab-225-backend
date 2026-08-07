from datetime import date
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

# Charte : orange Atelier.
_ORANGE = "F97316"
_WHITE = "FFFFFF"
_GREY = "F5F5F5"

_HEADERS = ["Référence", "Désignation", "Modèle Compatible", "Quantité totale"]
_QTY_COL = 4                      # colonne "Quantité totale"
_HEADER_ROW = 5                  # ligne des en-têtes
_DATA_START = _HEADER_ROW + 1    # première ligne de données

_thin = Side(style="thin", color="DDDDDD")
_border = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)


def _fmt_period(start_date: date | None, end_date: date | None) -> str:
    if start_date and end_date:
        return f"Période analysée : {start_date:%d/%m/%Y} au {end_date:%d/%m/%Y}"
    if start_date:
        return f"Période analysée : à partir du {start_date:%d/%m/%Y}"
    if end_date:
        return f"Période analysée : jusqu'au {end_date:%d/%m/%Y}"
    return "Période analysée : toutes périodes"


def _autosize(ws: Worksheet, *, n_cols: int, first_row: int, last_row: int) -> None:
    """Largeur auto SANS itérer ws.columns (qui plante sur les cellules fusionnées)."""
    for col in range(1, n_cols + 1):
        letter = get_column_letter(col)
        longest = len(_HEADERS[col - 1])
        for row in range(first_row, last_row + 1):
            value = ws.cell(row=row, column=col).value
            if value is not None:
                longest = max(longest, len(str(value)))
        ws.column_dimensions[letter].width = longest + 3


def build_parts_orders_workbook(
    rows: list[dict], *, start_date: date | None, end_date: date | None
) -> BytesIO:
    """rows : dicts issus de StatisticsService.parts_stats
    (reference, designation, models_label, total_quantity_ordered)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Commandes pièces"

    n_cols = len(_HEADERS)
    last_letter = get_column_letter(n_cols)

    # --- Titre ---
    ws.merge_cells(f"A1:{last_letter}1")
    title = ws["A1"]
    title.value = "Récapitulatif des commandes de pièces"
    title.font = Font(size=14, bold=True, color=_ORANGE)
    title.alignment = Alignment(horizontal="left", vertical="center")

    ws["A2"] = _fmt_period(start_date, end_date)
    ws["A2"].font = Font(italic=True, color="666666")
    ws["A3"] = f"Généré le : {date.today():%d/%m/%Y}"
    ws["A3"].font = Font(italic=True, color="666666")

    # --- En-têtes ---
    header_fill = PatternFill("solid", fgColor=_ORANGE)
    for col, label in enumerate(_HEADERS, start=1):
        c = ws.cell(row=_HEADER_ROW, column=col, value=label)
        c.font = Font(bold=True, color=_WHITE)
        c.fill = header_fill
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = _border

    # --- Données ---
    stripe = PatternFill("solid", fgColor=_GREY)
    row_idx = _DATA_START
    total_qty = 0
    for i, r in enumerate(rows):
        values = [
            r["reference"],
            r["designation"],
            r["models_label"],        # utilisation de la bonne clé
            r["total_quantity_ordered"],
        ]
        for col, value in enumerate(values, start=1):
            c = ws.cell(row=row_idx, column=col, value=value)
            c.border = _border
            if i % 2 == 1:
                c.fill = stripe
            if col == _QTY_COL:
                c.number_format = "#,##0"
                c.alignment = Alignment(horizontal="right")
        total_qty += int(r["total_quantity_ordered"])
        row_idx += 1

    # --- Ligne TOTAL ---
    if rows:
        label_cell = ws.cell(row=row_idx, column=_QTY_COL - 1, value="TOTAL")
        label_cell.font = Font(bold=True)
        label_cell.alignment = Alignment(horizontal="right")
        total_cell = ws.cell(row=row_idx, column=_QTY_COL, value=total_qty)
        total_cell.font = Font(bold=True)
        total_cell.number_format = "#,##0"
        total_cell.alignment = Alignment(horizontal="right")
        total_cell.fill = PatternFill("solid", fgColor="FFEAD5")  # orange très clair

    # --- Finitions ---
    _autosize(ws, n_cols=n_cols, first_row=_HEADER_ROW, last_row=row_idx)
    ws.freeze_panes = f"A{_DATA_START}"   # en-têtes figés au défilement
    ws.sheet_view.showGridLines = False

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer