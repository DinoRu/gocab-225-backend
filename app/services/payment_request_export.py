from collections import defaultdict
from datetime import date
from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

# --- Charte GOCAB ---
_NAVY = "1F3A5F"
_ORANGE = "F97316"
_WHITE = "FFFFFF"
_GREY = "F5F5F5"

# --- Couleurs de priorité (fond, texte) ---
_PRIORITY_STYLE: dict[str, tuple[str, str]] = {
    "low": ("EEF0F3", "6B7280"),       # gris
    "normal": ("E7F0FB", "2B6CB0"),    # bleu
    "high": ("FDF2E2", "B7791F"),      # ambre
    "urgent": ("FBEAE8", "C0392B"),    # rouge
}

# --- Couleurs de statut (fond, texte) ---
_STATUS_STYLE: dict[str, tuple[str, str]] = {
    "to_pay": ("FDF2E2", "B7791F"),    # à payer : ambre
    "paid": ("E6F4EC", "1E7D4F"),      # payé : vert
}

_PRIORITY_FR = {"low": "Basse", "normal": "Normale", "high": "Haute", "urgent": "Urgente"}
_STATUS_FR = {"to_pay": "À payer", "paid": "Payé"}

# Ordre d'affichage des priorités (urgent en premier).
_PRIORITY_ORDER = ["urgent", "high", "normal", "low"]

_HEADERS = ["N° demande", "Titre", "Fournisseur", "Réf. Odoo", "Priorité",
            "Date", "Statut", "Montant (FCFA)", "Lien Odoo"]
_AMOUNT_COL = 8
_LINK_COL = 9
_HEADER_ROW = 5
_DATA_START = 6

_thin = Side(style="thin", color="DDDDDD")
_border = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)


def build_payment_requests_workbook(rows: list[dict], *, title: str = "Demandes de paiement") -> BytesIO:
    wb = Workbook()

    # L'onglet dashboard est créé en premier (il apparaît en premier à l'ouverture).
    ws_dash = wb.active
    ws_dash.title = "Dashboard"
    _build_dashboard(ws_dash, rows, title=title)

    ws_detail = wb.create_sheet("Détail")
    _build_detail(ws_detail, rows, title=title)

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


# =========================================================================
# Onglet Détail
# =========================================================================
def _build_detail(ws: Worksheet, rows: list[dict], *, title: str) -> None:
    last = get_column_letter(len(_HEADERS))

    ws.merge_cells(f"A1:{last}1")
    c = ws["A1"]
    c.value = f"GOCAB 225 — {title}"
    c.font = Font(size=14, bold=True, color=_NAVY)
    c.alignment = Alignment(horizontal="left", vertical="center")

    ws["A3"] = f"Généré le : {date.today():%d/%m/%Y}"
    ws["A3"].font = Font(italic=True, color="666666")

    header_fill = PatternFill("solid", fgColor=_NAVY)
    for i, label in enumerate(_HEADERS, start=1):
        cell = ws.cell(row=_HEADER_ROW, column=i, value=label)
        cell.font = Font(bold=True, color=_WHITE)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = _border

    stripe = PatternFill("solid", fgColor=_GREY)
    total = Decimal("0")
    r = _DATA_START
    for idx, row in enumerate(rows):
        priority = row["priority"]
        status = row["status"]
        values = [
            row["request_number"],
            row["title"],
            row["supplier"]["name"],
            row.get("odoo_reference") or "",
            _PRIORITY_FR.get(priority, priority),
            _fmt_date(row["request_date"]),
            _STATUS_FR.get(status, status),
            float(row["amount"]),
            row.get("link") or "",
        ]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=r, column=col, value=value)
            cell.border = _border
            # Fond zébré léger (sauf sur les cellules qui portent leur propre couleur).
            if idx % 2 == 1 and col not in (5, 7):
                cell.fill = stripe

            if col == _AMOUNT_COL:
                cell.number_format = "#,##0"
                cell.alignment = Alignment(horizontal="right")

            # Colonne Priorité : fond + texte colorés selon le niveau.
            if col == 5:
                bg, fg = _PRIORITY_STYLE.get(priority, ("FFFFFF", "000000"))
                cell.fill = PatternFill("solid", fgColor=bg)
                cell.font = Font(bold=True, color=fg)
                cell.alignment = Alignment(horizontal="center")

            # Colonne Statut : fond + texte colorés.
            if col == 7:
                bg, fg = _STATUS_STYLE.get(status, ("FFFFFF", "000000"))
                cell.fill = PatternFill("solid", fgColor=bg)
                cell.font = Font(bold=True, color=fg)
                cell.alignment = Alignment(horizontal="center")

            # Colonne Lien : hyperlien cliquable.
            if col == _LINK_COL and value:
                cell.hyperlink = value
                cell.value = "Ouvrir dans Odoo"
                cell.font = Font(color="0563C1", underline="single")

        total += Decimal(str(row["amount"]))
        r += 1

    # Ligne TOTAL
    if rows:
        label = ws.cell(row=r, column=_AMOUNT_COL - 1, value="TOTAL")
        label.font = Font(bold=True)
        label.alignment = Alignment(horizontal="right")
        tot = ws.cell(row=r, column=_AMOUNT_COL, value=float(total))
        tot.font = Font(bold=True)
        tot.number_format = "#,##0"
        tot.fill = PatternFill("solid", fgColor="FFEAD5")

    widths = [14, 34, 22, 16, 12, 12, 12, 16, 22]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.freeze_panes = f"A{_DATA_START}"
    ws.sheet_view.showGridLines = False


# =========================================================================
# Onglet Dashboard (récapitulatif)
# =========================================================================
def _build_dashboard(ws: Worksheet, rows: list[dict], *, title: str) -> None:
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 16
    ws.column_dimensions["D"].width = 20
    ws.column_dimensions["E"].width = 4

    # Titre
    ws.merge_cells("B1:D1")
    t = ws["B1"]
    t.value = f"GOCAB 225 — {title}"
    t.font = Font(size=15, bold=True, color=_NAVY)
    ws["B2"] = f"Récapitulatif généré le {date.today():%d/%m/%Y}"
    ws["B2"].font = Font(italic=True, color="666666")

    # --- Agrégations (en Python : le dataset est déjà chargé et petit) ---
    total_amount = Decimal("0")
    by_status_count: dict[str, int] = defaultdict(int)
    by_status_amount: dict[str, Decimal] = defaultdict(lambda: Decimal("0"))
    by_priority_count: dict[str, int] = defaultdict(int)
    by_priority_amount: dict[str, Decimal] = defaultdict(lambda: Decimal("0"))

    for row in rows:
        amount = Decimal(str(row["amount"]))
        total_amount += amount
        by_status_count[row["status"]] += 1
        by_status_amount[row["status"]] += amount
        by_priority_count[row["priority"]] += 1
        by_priority_amount[row["priority"]] += amount

    # --- Cartes KPI (ligne 4) ---
    _kpi(ws, "B4", "Nombre de demandes", str(len(rows)), _NAVY)
    _kpi(ws, "C4", "Montant total", _fmt_money(total_amount), _ORANGE)

    # --- Bloc "Par statut" ---
    row_i = 8
    _section_title(ws, f"B{row_i}", "Répartition par statut")
    row_i += 1
    _table_header(ws, row_i, ["Statut", "Demandes", "Montant (FCFA)"])
    row_i += 1
    for status in ("to_pay", "paid"):
        bg, fg = _STATUS_STYLE.get(status, ("FFFFFF", "000000"))
        _row_cell(ws, row_i, 2, _STATUS_FR[status], fill=bg, color=fg, bold=True, center=True)
        _row_cell(ws, row_i, 3, by_status_count.get(status, 0), num=True)
        _row_cell(ws, row_i, 4, float(by_status_amount.get(status, Decimal("0"))), money=True)
        row_i += 1
    # total
    _row_cell(ws, row_i, 2, "Total", bold=True)
    _row_cell(ws, row_i, 3, len(rows), num=True, bold=True)
    _row_cell(ws, row_i, 4, float(total_amount), money=True, bold=True)

    # --- Bloc "Par priorité" ---
    row_i += 3
    _section_title(ws, f"B{row_i}", "Répartition par priorité")
    row_i += 1
    _table_header(ws, row_i, ["Priorité", "Demandes", "Montant (FCFA)"])
    row_i += 1
    for priority in _PRIORITY_ORDER:
        if by_priority_count.get(priority, 0) == 0:
            continue  # on n'affiche que les priorités présentes
        bg, fg = _PRIORITY_STYLE.get(priority, ("FFFFFF", "000000"))
        _row_cell(ws, row_i, 2, _PRIORITY_FR[priority], fill=bg, color=fg, bold=True, center=True)
        _row_cell(ws, row_i, 3, by_priority_count.get(priority, 0), num=True)
        _row_cell(ws, row_i, 4, float(by_priority_amount.get(priority, Decimal("0"))), money=True)
        row_i += 1


# =========================================================================
# Helpers de mise en forme (dashboard)
# =========================================================================
def _kpi(ws: Worksheet, anchor: str, label: str, value: str, color: str) -> None:
    col = anchor[0]
    row = int(anchor[1:])
    lab = ws[f"{col}{row}"]
    lab.value = label
    lab.font = Font(size=9, bold=True, color="6B7280")
    val = ws[f"{col}{row + 1}"]
    val.value = value
    val.font = Font(size=16, bold=True, color=color)


def _section_title(ws: Worksheet, anchor: str, text: str) -> None:
    cell = ws[anchor]
    cell.value = text
    cell.font = Font(size=12, bold=True, color=_NAVY)


def _table_header(ws: Worksheet, row: int, labels: list[str]) -> None:
    fill = PatternFill("solid", fgColor=_NAVY)
    for i, label in enumerate(labels):
        cell = ws.cell(row=row, column=2 + i, value=label)
        cell.font = Font(bold=True, color=_WHITE, size=10)
        cell.fill = fill
        cell.alignment = Alignment(horizontal="left" if i == 0 else "right")
        cell.border = _border


def _row_cell(
    ws: Worksheet, row: int, col: int, value,
    *, fill: str | None = None, color: str | None = None,
    bold: bool = False, center: bool = False, num: bool = False, money: bool = False,
) -> None:
    cell = ws.cell(row=row, column=col, value=value)
    cell.border = _border
    if fill:
        cell.fill = PatternFill("solid", fgColor=fill)
    cell.font = Font(bold=bold, color=color or "000000")
    if money:
        cell.number_format = "#,##0"
        cell.alignment = Alignment(horizontal="right")
    elif num:
        cell.number_format = "#,##0"
        cell.alignment = Alignment(horizontal="right")
    elif center:
        cell.alignment = Alignment(horizontal="center")


def _fmt_money(value: Decimal) -> str:
    return f"{value:,.0f}".replace(",", " ") + " FCFA"


def _fmt_date(value) -> str:
    if isinstance(value, date):
        return value.strftime("%d/%m/%Y")
    return str(value)