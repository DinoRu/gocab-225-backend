from datetime import date
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

_NAVY = "1F3A5F"
_WHITE = "FFFFFF"
_GREY = "F5F5F5"

_STATUS_STYLE = {
    "open": ("FDF2E2", "B7791F"),
    "in_progress": ("E7F0FB", "2B6CB0"),
    "fulfilled": ("E6F4EC", "1E7D4F"),
}
_STATUS_FR = {"open": "Ouvert", "in_progress": "En cours", "fulfilled": "Traité"}

_thin = Side(style="thin", color="DDDDDD")
_border = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)


def build_supply_requests_workbook(rows: list[dict], *, status_filter=None) -> BytesIO:
    """Une ligne par (besoin × pièce) — vue à plat pratique pour traiter les commandes."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Besoins"

    title = "Besoins d'approvisionnement"
    if status_filter:
        title += f" — {_STATUS_FR.get(status_filter, status_filter)}"

    ws.merge_cells("A1:F1")
    ws["A1"] = f"GOCAB 225 — {title}"
    ws["A1"].font = Font(size=14, bold=True, color=_NAVY)
    ws["A3"] = f"Généré le : {date.today():%d/%m/%Y}"
    ws["A3"].font = Font(italic=True, color="666666")

    headers = ["N° besoin", "Date", "Statut", "Référence", "Désignation", "Quantité"]
    header_row = 5
    fill = PatternFill("solid", fgColor=_NAVY)
    for i, label in enumerate(headers, start=1):
        c = ws.cell(row=header_row, column=i, value=label)
        c.font = Font(bold=True, color=_WHITE)
        c.fill = fill
        c.alignment = Alignment(horizontal="center")
        c.border = _border

    r = header_row + 1
    stripe = PatternFill("solid", fgColor=_GREY)
    even = False
    for sr in rows:
        status = sr["status"]
        bg, fg = _STATUS_STYLE.get(status, ("FFFFFF", "000000"))
        for it in sr["items"]:
            values = [
                sr["sr_number"],
                _fmt(sr["request_date"]),
                _STATUS_FR.get(status, status),
                it["reference"],
                it["designation"],
                it["quantity"],
            ]
            for col, val in enumerate(values, start=1):
                cell = ws.cell(row=r, column=col, value=val)
                cell.border = _border
                if even:
                    cell.fill = stripe
                if col == 3:  # statut coloré
                    cell.fill = PatternFill("solid", fgColor=bg)
                    cell.font = Font(bold=True, color=fg)
                    cell.alignment = Alignment(horizontal="center")
                if col == 6:
                    cell.alignment = Alignment(horizontal="right")
                    cell.number_format = "#,##0"
            r += 1
        even = not even

    widths = [14, 12, 12, 16, 34, 10]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = f"A{header_row + 1}"
    ws.sheet_view.showGridLines = False

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def _fmt(value) -> str:
    if isinstance(value, date):
        return value.strftime("%d/%m/%Y")
    return str(value)