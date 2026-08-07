from datetime import date
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

_NAVY = "1F3A5F"
_ORANGE = "F97316"
_WHITE = "FFFFFF"
_RED = "C0392B"
_GREY = "F5F5F5"

_HEADERS = ["Référence", "Désignation", "Stock précédent", "Entrées",
            "Stock compté", "Sorties", "Anomalie"]
_HEADER_ROW = 5
_DATA_START = 6

_thin = Side(style="thin", color="DDDDDD")
_border = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)


def build_inventory_workbook(count: dict) -> BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = "Inventaire"

    last = get_column_letter(len(_HEADERS))
    ws.merge_cells(f"A1:{last}1")
    c = ws["A1"]
    c.value = f"GOCAB 225 — Comptage {count['count_number']}"
    c.font = Font(size=14, bold=True, color=_NAVY)

    ws["A2"] = f"Date du comptage : {_fmt(count['count_date'])}"
    ws["A2"].font = Font(italic=True, color="666666")
    ws["A3"] = f"Généré le : {date.today():%d/%m/%Y}"
    ws["A3"].font = Font(italic=True, color="666666")

    fill = PatternFill("solid", fgColor=_NAVY)
    for i, label in enumerate(_HEADERS, start=1):
        cell = ws.cell(row=_HEADER_ROW, column=i, value=label)
        cell.font = Font(bold=True, color=_WHITE)
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = _border

    stripe = PatternFill("solid", fgColor=_GREY)
    r = _DATA_START
    for idx, it in enumerate(count["items"]):
        first = it["previous_quantity"] is None
        values = [
            it["reference"],
            it["designation"],
            "—" if first else it["previous_quantity"],
            "—" if first else it["entries_between"],
            it["counted_quantity"],
            "—" if it["outflow"] is None else it["outflow"],
            "⚠ Incohérence" if it["anomaly"] else "",
        ]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=r, column=col, value=value)
            cell.border = _border
            if idx % 2 == 1:
                cell.fill = stripe
            if col in (3, 4, 5, 6):
                cell.alignment = Alignment(horizontal="right")
                if isinstance(value, int):
                    cell.number_format = "#,##0"
            if col == 6 and it["anomaly"]:
                cell.font = Font(bold=True, color=_RED)
            if col == 7 and it["anomaly"]:
                cell.font = Font(bold=True, color=_RED)
        r += 1

    widths = [16, 30, 15, 12, 14, 12, 16]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.freeze_panes = f"A{_DATA_START}"
    ws.sheet_view.showGridLines = False

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def _fmt(value) -> str:
    if isinstance(value, date):
        return value.strftime("%d/%m/%Y")
    return str(value)